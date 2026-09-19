from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

W, H = 13.333, 7.5
DARK = RGBColor(18, 11, 16)
MAUVE = RGBColor(67, 40, 58)
MIST = RGBColor(238, 220, 230)
GOLD = RGBColor(245, 196, 107)
IVORY = RGBColor(252, 251, 249)
INK = RGBColor(36, 28, 23)
MUTED = RGBColor(120, 107, 98)
GREEN = RGBColor(45, 138, 67)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)


def add_shape(slide, kind, x, y, w, h, color, rounded=False, transparency=0, rotation=0):
    item = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = color
    item.fill.transparency = transparency
    item.line.fill.background()
    item.rotation = rotation
    return item


def add_text(slide, value, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Aptos"):
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = item.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return item


def add_lily(slide, x, y, scale=1):
    for offset_x, offset_y in [(-28, 0), (-14, -22), (0, -32), (14, -22), (28, 0)]:
        add_shape(slide, MSO_SHAPE.OVAL, x + (offset_x + 34) * scale, y + (offset_y + 30) * scale, 32 * scale, 70 * scale, RGBColor(185, 141, 169), True, 18, offset_x / 2)
    add_shape(slide, MSO_SHAPE.OVAL, x + 30 * scale, y + 27 * scale, 40 * scale, 40 * scale, GOLD, True, 5)
    add_shape(slide, MSO_SHAPE.RECTANGLE, x + 48 * scale, y + 62 * scale, 4 * scale, 145 * scale, RGBColor(90, 61, 82), rotation=-2)


def add_footer(slide, number):
    add_text(slide, "LEI  /  Hackathon pitch", 0.72, 7.08, 4, 0.2, 9, MIST)
    add_text(slide, f"0{number}", 12.1, 7.08, 0.5, 0.2, 9, GOLD, True, PP_ALIGN.RIGHT)


def new_slide(dark=True):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, H, DARK if dark else IVORY)
    return slide


def add_heading(slide, kicker, title, number):
    add_text(slide, kicker.upper(), 0.72, 0.42, 5.5, 0.25, 10, GOLD, True)
    add_text(slide, title, 0.72, 0.78, 11.6, 0.75, 29, MIST, True, font="Aptos Display")
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.72, 1.68, 1.05, 0.04, GOLD)
    add_footer(slide, number)


# Slide 1: The hook
slide = new_slide()
add_shape(slide, MSO_SHAPE.OVAL, 8.6, -1.0, 6.2, 7.0, MAUVE, True, 45)
add_shape(slide, MSO_SHAPE.OVAL, 10.4, 1.1, 3.7, 4.8, GOLD, True, 75)
add_lily(slide, 9.55, 2.55, 1.65)
add_text(slide, "LÈI", 0.85, 0.82, 3.5, 0.45, 18, GOLD, True)
add_text(slide, "The internet was built\nfor engagement.\nWomen need it built\nfor safety.", 0.85, 1.8, 7.3, 2.4, 34, MIST, True, font="Aptos Display")
add_text(slide, "Every day, women navigate digital spaces that are fragmented, unsafe, or optimized for engagement rather than safety.", 0.9, 4.55, 7.2, 0.85, 18, MIST)
add_text(slide, "Meet Lèi - a gated, trauma-informed digital ecosystem built exclusively for women to connect, discover, and grow safely.", 0.9, 5.72, 7.3, 0.7, 16, GOLD, True)
add_footer(slide, 1)

# Slide 2: Education and career growth
slide = new_slide()
add_heading(slide, "02  /  Track alignment", "Empowerment becomes tangible when support leads somewhere.", 2)
add_text(slide, "Lèi goes beyond a chat app: it turns knowledge, opportunity, and immediate help into an accessible daily experience.", 0.82, 2.0, 11.6, 0.48, 19, MIST)
features = [
    ("EDUCATION", "Discord-style hubs", "Women in STEM, higher education, grad research, and first-gen communities share knowledge and opportunities."),
    ("MUTUAL AID", "Campus SOS", "A student at GSU can request a sanitary pad, form a study circle, or reach nearby peers in minutes."),
    ("CAREER", "Lèi AI Concierge", "Resume feedback, career coaching, wellness grounding, and a clear next action in one private conversation."),
]
for index, (tag, title, body) in enumerate(features):
    x = 0.82 + index * 4.1
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.95, 3.55, 2.65, RGBColor(42, 28, 38), True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, 2.95, 0.1, 2.65, GOLD if index == 1 else RGBColor(185, 141, 169))
    add_text(slide, tag, x + 0.3, 3.25, 2.75, 0.24, 10, GOLD, True)
    add_text(slide, title, x + 0.3, 3.68, 2.9, 0.5, 22, MIST, True, font="Aptos Display")
    add_text(slide, body, x + 0.3, 4.45, 2.9, 0.85, 14, MIST)
add_text(slide, "Track outcome: more access to learning, more career agency, more women able to help one another.", 0.82, 6.25, 11.6, 0.4, 18, GOLD, True)

# Slide 3: Meaningful connections and safety
slide = new_slide(False)
add_heading(slide, "03  /  Meaningful connection", "Digital connection should translate into real-world safety.", 3)
add_text(slide, "Lèi gives connection context: how I feel, where I am, and what kind of support I need today.", 0.82, 2.0, 11.6, 0.48, 19, INK)
features = [
    ("WELLBEING", "What part of HER are you today?", "Teams-style status lets members signal their state without changing what the algorithm shows them."),
    ("LOCAL", "USA City Meetups", "Filter by city and state to build trusted friendships in verified public settings, like NFS in Atlanta."),
    ("SAFETY", "Beacons and legal support", "Safety beacons, rights guides, emergency hotlines, and support timers connect digital care to physical reality."),
]
for index, (tag, title, body) in enumerate(features):
    x = 0.82 + index * 4.1
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.95, 3.55, 2.65, RGBColor(255, 255, 255), True)
    add_text(slide, tag, x + 0.28, 3.25, 2.9, 0.25, 10, MAUVE, True)
    add_text(slide, title, x + 0.28, 3.67, 2.9, 0.62, 20, INK, True, font="Aptos Display")
    add_text(slide, body, x + 0.28, 4.5, 2.92, 0.78, 14, MUTED)
add_text(slide, "Meaningful connection is not a metric. It is a person knowing where to turn next.", 0.82, 6.25, 11.6, 0.4, 18, MAUVE, True)

# Slide 4: Zero-trust architecture
slide = new_slide(False)
add_heading(slide, "04  /  Technical edge", "Trust is not a setting. It is the architecture.", 4)
add_text(slide, "Lèi treats safety as the foundation that makes education, career growth, and connection possible.", 0.82, 2.0, 11.6, 0.48, 19, INK)
architecture = [
    ("01", "Credential uniqueness", "Phone and email indexing reduce burner accounts and duplicate signups."),
    ("02", "Multi-factor onboarding", "Camera liveness checks through MediaDevices API plus real 6-digit SMS verification in production."),
    ("03", "Spatial access control", "Members choose a primary sanctuary space; cross-space access is locked behind an appeal workflow."),
]
for index, (number, title, body) in enumerate(architecture):
    y = 2.95 + index * 1.05
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.82, y, 11.4, 0.82, RGBColor(255, 255, 255), True)
    add_text(slide, number, 1.1, y + 0.18, 0.55, 0.35, 15, MAUVE, True)
    add_text(slide, title, 1.9, y + 0.15, 3.0, 0.3, 17, INK, True)
    add_text(slide, body, 5.0, y + 0.12, 6.65, 0.4, 14, MUTED)
add_text(slide, "The technical edge is not surveillance. It is intentional boundaries that protect agency.", 0.82, 6.35, 11.6, 0.4, 18, GREEN, True)

# Slide 5: Conclusion
slide = new_slide()
add_shape(slide, MSO_SHAPE.OVAL, 8.8, -0.8, 5.6, 7.2, MAUVE, True, 45)
add_lily(slide, 10.0, 2.0, 1.8)
add_text(slide, "LÈI", 0.85, 1.05, 3.0, 0.45, 18, GOLD, True)
add_text(slide, "Infrastructure for\nwomen's autonomy.", 0.85, 2.0, 7.5, 1.3, 37, MIST, True, font="Aptos Display")
add_text(slide, "When women are safe, supported, and connected, they change the world.", 0.9, 4.1, 7.4, 0.55, 22, GOLD, True)
add_text(slide, "Thank you.", 0.9, 5.35, 3.0, 0.45, 20, MIST)
add_text(slide, "Live demo: pkyeibrewu1.github.io/LEI/?developer=1", 0.9, 6.22, 7.5, 0.28, 11, MIST)
add_footer(slide, 5)

prs.save("presentation/lei-hackathon-pitch.pptx")
